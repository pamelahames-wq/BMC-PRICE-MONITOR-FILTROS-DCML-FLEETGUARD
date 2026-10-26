"""Descontos RAP por grupo de peça: margem planejada (Reestruturação de Margem 2026) vs margem real.

Uso:  python -m rap.ajuste_grupos   →  saida/rap/Ajuste_desconto_por_grupo.xlsx
Lê a aba "Resumo por grupo" da Reestruturacao_Margem_2026_FINAL.xlsx copiada em
cache/raw/resumo_margem_grupos.tsv e as NFs RAP do 2º sem/2026 (rap.analise).
A planilha da reestruturação desconta 13,33% de impostos; nas NFs RAP o imposto médio é ~20,5%.
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import analise as A
from .saidas import BOLD, INPUT, PCT, RS0, SAIDA, TITULO, _cab, _filtro, _fmt

ARQ = SAIDA / "Ajuste_desconto_por_grupo.xlsx"
IMP_PLANILHA = 0.1333   # imposto implícito nas colunas "Margem RAP" da reestruturação (ajuste com erro < 0,1 p.p.)
MIN_FAT = 5000          # abaixo disso usa custo/tabela planejado e imposto médio RAP


def _pct(s):
    return pd.to_numeric(s.str.rstrip("%"), errors="coerce") / 100


def dados():
    g = pd.read_csv(A.RAW / "resumo_margem_grupos.tsv", sep="\t", dtype=str)
    g.columns = ["n1", "cod", "grupo", "itens", "m_atual", "m_15", "dS", "dO", "dD", "mb", "mS", "mO", "mD"]
    for c in ["m_atual", "m_15", "dS", "dO", "dD", "mb", "mS", "mO", "mD"]:
        g[c] = _pct(g[c])
    g = g.dropna(subset=["dD"])
    g["r_plan"] = (1 - IMP_PLANILHA - g.mb) * (1 - A.DESC_BALCAO)

    r = A.rodar()
    _, _, _, _, pre, _ = A.carregar()
    precos = dict(zip(pre.CODPROD, pre.preco))
    v = r["itens"][r["itens"].sinal > 0].copy()
    v["tab"] = v.CODPROD.map(precos) * v.QTDNEG
    v = v[v.tab > 0]
    s = v.groupby("CODGRUPOPROD").agg(fat=("BRUTO", "sum"), imp=("IMPOSTOS", "sum"), tab=("tab", "sum"),
                                       cst=("custo_total", "sum"))
    imp_medio = s.imp.sum() / s.fat.sum()
    g = g.merge(s, left_on="cod", right_index=True, how="left")
    real = g.fat >= MIN_FAT
    g["imp"] = (g.imp / g.fat).where(real, imp_medio)
    g["r"] = (g.cst / g.tab).where(real, g.r_plan)
    g["fonte"] = np.where(real, "NFs RAP 2º sem", "planejado (pouca venda RAP)")
    g["fat"] = g.fat.fillna(0)
    return g.sort_values("fat", ascending=False), imp_medio


def gerar(piso=0.05):
    g, imp_medio = dados()
    wb = Workbook()
    ws = wb.active
    ws.title = "Desconto por grupo"
    ws["A1"] = "Desconto RAP por grupo de peça — margem planejada vs margem com o imposto real das NFs RAP"
    ws["A1"].font = TITULO
    ws["A2"] = "Piso de margem RAP"
    ws["A2"].font = BOLD
    _fmt(ws, 2, 3, piso, PCT, INPUT).fill = PatternFill("solid", fgColor="FFFF00")
    ws["D2"] = ("Margem na convenção da reestruturação: (preço − impostos − custo) ÷ preço. Desconto sugerido = "
                "menor entre o atual e o maior desconto que mantém a margem no piso (célula C2, editável).")
    ws["D2"].font = Font(name="Arial", italic=True)
    cols = ["CODGRUPOPROD", "Grupo", "Itens", "Faturamento RAP 2º sem (R$)", "Fonte custo/imposto",
            "Imposto ÷ NF", "Custo ÷ tabela",
            "Desc STD atual", "Desc OURO atual", "Desc DIAMANTE atual",
            "Margem DIAMANTE planilha (13,3% imp.)",
            "Margem real STD", "Margem real OURO", "Margem real DIAMANTE",
            "Desc STD sugerido", "Desc OURO sugerido", "Desc DIAMANTE sugerido",
            "Margem DIAMANTE no sugerido", "Muda?"]
    _cab(ws, 4, cols, [12, 34, 7, 15, 16, 9, 9, 9, 9, 10, 12, 9, 9, 10, 10, 10, 11, 11, 8])
    ws.row_dimensions[4].height = 58
    for i, x in enumerate(g.itertuples(), 5):
        _fmt(ws, i, 1, int(x.cod) if str(x.cod).isdigit() else x.cod)
        _fmt(ws, i, 2, x.grupo)
        _fmt(ws, i, 3, int(x.itens))
        _fmt(ws, i, 4, round(x.fat, 2), RS0)
        _fmt(ws, i, 5, x.fonte)
        _fmt(ws, i, 6, round(x.imp, 4), PCT, INPUT)
        _fmt(ws, i, 7, round(x.r, 4), PCT, INPUT)
        for j, d in enumerate([x.dS, x.dO, x.dD], 8):
            _fmt(ws, i, j, d, PCT, INPUT)
        _fmt(ws, i, 11, x.mD, PCT)
        for j, col in enumerate("HIJ", 12):           # margem real com o desconto atual
            _fmt(ws, i, j, f"=1-$F{i}-$G{i}/(1-{col}{i})", PCT)
        for j, col in enumerate("HIJ", 15):           # desconto sugerido respeitando o piso
            _fmt(ws, i, j, f"=MAX(0,MIN({col}{i},ROUNDDOWN(1-$G{i}/(1-$F{i}-$C$2),2)))", PCT)
        _fmt(ws, i, 18, f"=1-$F{i}-$G{i}/(1-Q{i})", PCT)
        _fmt(ws, i, 19, f'=IF(OR(O{i}<H{i},P{i}<I{i},Q{i}<J{i}),"SIM","")', None, BOLD)
    _filtro(ws, 4, len(g) + 4, len(cols))
    for row in ws.iter_rows(min_row=5, max_row=len(g) + 4, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=False)
    n = len(g) + 6
    ws.cell(n, 1, f"Imposto médio nas NFs RAP (2º sem/2026): {imp_medio:.1%}; a planilha da reestruturação usa "
                  f"13,3%, que é o nível do Balcão. Grupos com menos de R$ {MIN_FAT:,} vendidos para RAP usam o "
                  "custo/tabela planejado e o imposto médio.".replace(",", ".")).font = Font(name="Arial", italic=True)
    wb.save(ARQ)
    return ARQ


if __name__ == "__main__":
    print(gerar())
