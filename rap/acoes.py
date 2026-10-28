"""Listas de ação da reestruturação RAP: ligações, listas por consultor e encerramentos.

Uso:  python -m rap.acoes
Gera em saida/rap/:
  Acoes_RAP_2S2026.xlsx           — todas as listas
  consultores/Lista_<CONSULTOR>.xlsx — só os clientes de cada consultor
Telefones: cache/raw/hubspot_rap_empresas.json (empresas do HubSpot por CNPJ) e CONTATO da d_Clientes.
"""
import json
import shutil

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from campanha import norm
from . import analise as A
from .saidas import BOLD, DATA, INPUT, INT, PCT, RS, TITULO, _cab, _filtro, _fmt

SAIDA = A.RAIZ / "saida" / "rap"
ARQ = SAIDA / "Acoes_RAP_2S2026.xlsx"
HUB_URL = "https://app.hubspot.com/contacts/9472165/record/0-2/{}"
AMARELO = PatternFill("solid", fgColor="FFF2CC")


def _fone(d):
    d = norm.so_digitos(d)
    while d.startswith("55") and len(d) > 11:
        d = d[2:]
    d = d.lstrip("0")
    if len(d) == 10:
        return f"({d[:2]}) {d[2:6]}-{d[6:]}"
    if len(d) == 11:
        return f"({d[:2]}) {d[2:7]}-{d[7:]}"
    return ""


def contatos(c):
    arq = A.RAW / "hubspot_rap_empresas.json"
    hub = json.loads(arq.read_text()) if arq.exists() else {}
    fones, links = [], []
    for x in c.itertuples():
        achados = []
        for h in hub.get(x.CNPJ, []):
            for k in ("phone", "tel_cnpj", "tel_bkp"):
                f = _fone(h.get(k) or "")
                if f and f not in achados:
                    achados.append(f)
        for t in norm.telefones(x.CONTATO):
            f = _fone(t)
            if f and f not in achados:
                achados.append(f)
        fones.append(" / ".join(achados[:2]))
        links.append(HUB_URL.format(hub[x.CNPJ][0]["id"]) if hub.get(x.CNPJ) else "")
    return fones, links


def recorrencia(fat):
    v = fat[(fat.TIPMOV == "V") & (fat.DT >= A.INICIO) & (fat.DT <= A.FIM)]
    g = v.groupby("CODPARC")
    return g.NUNOTA.nunique(), g.DT.max()


def listas():
    r = A.rodar()
    _, fat, _, _, _, _ = A.carregar()
    c = r["classificacao"].copy()
    nfs, ult = recorrencia(fat)
    c["nfs"] = c.CODPARC.map(nfs).fillna(0).astype(int)
    c["ultima"] = c.CODPARC.map(ult)
    c["fone"], c["hub"] = contatos(c)
    c["consultor"] = c.consultor.str.replace(" (ELISASILVA)", "", regex=False)
    desh = c.balde == "DESHOMOLOGAR"
    return {
        "ligar": c[desh & (c.fat > 0)].sort_values("fat", ascending=False),
        "ouro": c[(c.balde == "REBAIXAR") & (c.categoria == "Ouro")].sort_values(["consultor", "pct"],
                                                                               ascending=[True, False]),
        "diamante": c[(c.balde == "REBAIXAR") & (c.categoria == "Diamante")].sort_values(["consultor", "pct"],
                                                                                       ascending=[True, False]),
        "encerrar": c[desh & (c.fat <= 0)].sort_values(["consultor", "categoria", "NOMEPARC"]),
    }


def aba_meta(wb, nome, d, titulo, sub, minimo_txt, com_resultado=False):
    ws = wb.create_sheet(nome)
    ws["A1"] = titulo
    ws["A1"].font = TITULO
    ws["A2"] = sub
    ws["A2"].font = Font(name="Arial", italic=True)
    cols = ["Consultor", "RAP", "CNPJ", "Telefone", "E-mail", "Categoria atual",
            "Faturado no semestre (R$)", f"Mínimo {minimo_txt} (R$)", "Quanto falta (R$)", "% do mínimo",
            "Nº de NFs no semestre", "Última compra", "Comprou no 1º sem/2026 (R$)", "HubSpot"]
    if com_resultado:
        cols += ["Resultado da ligação", "Data do contato"]
    _cab(ws, 4, cols, [13, 38, 17, 26, 30, 11, 15, 13, 14, 10, 10, 12, 16, 12, 34, 13][:len(cols)])
    for i, x in enumerate(d.itertuples(), 5):
        _fmt(ws, i, 1, x.consultor)
        _fmt(ws, i, 2, x.NOMEPARC)
        _fmt(ws, i, 3, x.CNPJ)
        _fmt(ws, i, 4, x.fone)
        _fmt(ws, i, 5, x.EMAIL if isinstance(x.EMAIL, str) else "")
        _fmt(ws, i, 6, x.categoria)
        _fmt(ws, i, 7, x.fat, RS, INPUT)
        _fmt(ws, i, 8, x.minimo, RS, INPUT)
        _fmt(ws, i, 9, f"=MAX(0,H{i}-G{i})", RS, BOLD)
        _fmt(ws, i, 10, f"=IF(H{i}>0,G{i}/H{i},0)", PCT)
        _fmt(ws, i, 11, x.nfs, INT)
        _fmt(ws, i, 12, x.ultima.to_pydatetime() if pd.notna(x.ultima) else None, DATA)
        _fmt(ws, i, 13, x.fat_1s, RS)
        if x.hub:
            c = _fmt(ws, i, 14, "abrir")
            c.hyperlink = x.hub
            c.font = Font(name="Arial", color="0563C1", underline="single")
        if com_resultado:
            ws.cell(i, 15).fill = AMARELO
            ws.cell(i, 16).fill = AMARELO
    n = len(d) + 5
    _fmt(ws, n, 2, "Total", None, BOLD)
    for col, fmt in (("G", RS), ("I", RS)):
        _fmt(ws, n, " ABCDEFGHI".index(col), f"=SUM({col}5:{col}{n - 1})", fmt, BOLD)
    _filtro(ws, 4, n - 1, len(cols))
    ws.freeze_panes = "C5"
    return ws


def aba_encerrar(wb, d):
    ws = wb.create_sheet("Encerramento (385)")
    ws["A1"] = f"Encerramento de contrato — {len(d)} RAPs sem faturamento no 2º semestre/2026 (01/07 a 16/10)"
    ws["A1"].font = TITULO
    ws["A2"] = ("Base para mala direta do e-mail e da carta de encerramento. Revisar antes as que compraram no "
                "1º semestre (coluna H) e as sem e-mail. Enviar só depois da aprovação.")
    ws["A2"].font = Font(name="Arial", italic=True)
    cols = ["CONSULTOR", "RAZAO_SOCIAL", "CNPJ", "E-MAIL", "CATEGORIA", "Cidade", "UF",
            "Comprou no 1º sem/2026 (R$)", "Revisar antes?", "Enviado em"]
    _cab(ws, 4, cols, [13, 40, 17, 34, 11, 20, 5, 16, 30, 12])
    rep = d.EMAIL.str.lower().str.strip().duplicated(keep="first")
    for i, (x, dup) in enumerate(zip(d.itertuples(), rep), 5):
        for j, v in enumerate([x.consultor, x.NOMEPARC, x.CNPJ, x.EMAIL if isinstance(x.EMAIL, str) else "",
                               x.categoria, x.CIDADE, x.UF], 1):
            _fmt(ws, i, j, v)
        _fmt(ws, i, 8, x.fat_1s, RS)
        motivo = []
        if x.fat_1s > 0:
            motivo.append("comprou no 1º sem")
        if not isinstance(x.EMAIL, str):
            motivo.append("sem e-mail")
        if x.fat < 0:
            motivo.append("devolução > compras")
        if dup:
            motivo.append("e-mail já aparece acima (filial)")
        _fmt(ws, i, 9, "; ".join(motivo))
        ws.cell(i, 10).fill = AMARELO
    _filtro(ws, 4, len(d) + 4, len(cols))
    ws.freeze_panes = "C5"


def aba_resumo(wb, L):
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Reestruturação RAP — listas de ação (NFs de 01/07 a 16/10/2026)"
    ws["A1"].font = TITULO
    _cab(ws, 3, ["Consultor", "Ligar (Standard abaixo de 70%)", "Ouro → Standard", "Diamante → Ouro",
                 "Encerramento", "Total"], [14, 18, 16, 16, 14, 10])
    cons = sorted(set().union(*[set(d.consultor) for d in L.values()]))
    for i, k in enumerate(cons + ["Total"], 4):
        _fmt(ws, i, 1, k, None, BOLD if k == "Total" else None)
        for j, chave in enumerate(["ligar", "ouro", "diamante", "encerrar"], 2):
            v = len(L[chave]) if k == "Total" else int((L[chave].consultor == k).sum())
            _fmt(ws, i, j, v, INT, BOLD if k == "Total" else None)
        _fmt(ws, i, 6, f"=SUM(B{i}:E{i})", INT, BOLD)
    n = len(cons) + 7
    notas = [
        "Standard: para manter os 20% de desconto precisa faturar no mínimo R$ 5.000 no semestre, com recorrência "
        "de compra. As 12 da aba 'Ligar' faturaram algo, mas abaixo de 70% do mínimo — ligação da Pamela.",
        "Ouro → Standard e Diamante → Ouro: 'Quanto falta' = mínimo da categoria atual − faturado até 16/10. "
        "O semestre fecha em 31/12: quem completar o mínimo até lá mantém a categoria.",
        "Encerramento: só depois da aprovação. Modelo de e-mail e carta no arquivo Encerramento_RAP_modelo.docx.",
        "Telefones: HubSpot (empresa pelo CNPJ) e campo CONTATO do Sankhya.",
    ]
    for k, t in enumerate(notas):
        ws.cell(n + k, 1, "• " + t).font = Font(name="Arial")


def gerar():
    L = listas()
    SAIDA.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    aba_resumo(wb, L)
    aba_meta(wb, "Ligar - Standard (12)", L["ligar"],
             "Ligar — 12 RAPs Standard que faturaram abaixo de 70% do mínimo (responsável: Pamela)",
             "Para manter 20% de desconto: mínimo de R$ 5.000 no semestre, com recorrência de compra.",
             "Standard", com_resultado=True)
    aba_meta(wb, "Ouro → Standard", L["ouro"],
             f"Ouro → Standard — {len(L['ouro'])} RAPs abaixo de 70% do mínimo de R$ 40.000",
             "Quanto falta para continuar Ouro até 31/12. Lista direcionada por consultor (filtre a coluna A).",
             "Ouro", com_resultado=True)
    aba_meta(wb, "Diamante → Ouro", L["diamante"],
             f"Diamante → Ouro — {len(L['diamante'])} RAPs abaixo de 70% do mínimo de R$ 100.000",
             "Quanto falta para continuar Diamante até 31/12. Lista direcionada por consultor (filtre a coluna A).",
             "Diamante", com_resultado=True)
    aba_encerrar(wb, L["encerrar"])
    wb.save(ARQ)

    pasta = SAIDA / "consultores"
    if pasta.exists():
        shutil.rmtree(pasta)
    pasta.mkdir()
    for k in sorted(set(L["ouro"].consultor) | set(L["diamante"].consultor) | set(L["encerrar"].consultor)):
        wb = Workbook()
        ws = wb.active
        ws.title = "Leia-me"
        ws["A1"] = f"Lista de clientes RAP — {k}"
        ws["A1"].font = TITULO
        ws["A3"] = ("Faturamento de 01/07 a 16/10/2026. 'Quanto falta' = valor para manter a categoria atual até "
                    "31/12. Registre o resultado de cada contato nas colunas amarelas.")
        ws["A3"].font = Font(name="Arial")
        for chave, nome, cat, minimo in (("diamante", "Diamante → Ouro", "Diamante", "Diamante"),
                                         ("ouro", "Ouro → Standard", "Ouro", "Ouro")):
            d = L[chave][L[chave].consultor == k]
            if len(d):
                aba_meta(wb, nome, d, f"{nome} — {len(d)} RAPs de {k}",
                         f"Abaixo de 70% do mínimo da categoria {cat}.", minimo, com_resultado=True)
        d = L["encerrar"][L["encerrar"].consultor == k]
        if len(d):
            aba_encerrar(wb, d)
            wb[wb.sheetnames[-1]].title = f"Encerramento ({len(d)})"
        wb.save(pasta / f"Lista_{k}.xlsx")
    return L


if __name__ == "__main__":
    L = gerar()
    for k, d in L.items():
        print(k, len(d), "com telefone:", (d.fone != "").sum(), "com e-mail:", d.EMAIL.notna().sum())
