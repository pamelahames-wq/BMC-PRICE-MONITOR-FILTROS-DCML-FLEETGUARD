"""Auditoria das planilhas por consultor da campanha GPD sem depender das bases do Sankhya.

Uso:  python -m campanha.auditar <pasta> [saida.xlsx]
A pasta tem uma subpasta por consultor com Repeticao_X e Novos_Clientes_X (.xlsx ou .tsv).
Cliente = CNPJ/CPF completo: nomes iguais com documentos diferentes não são juntados, só apontados.
"""
import re
import sys
from pathlib import Path

import pandas as pd

from . import norm
from .gerar import ES_NORTE, MG_DDD, RAIZ_EXCLUIR, UF_MAPA, _sem_acento

INI, FIM = pd.Timestamp("2026-01-01"), pd.Timestamp("2026-06-30")
REVENDA = re.compile(r"\bPECAS?\b|\bPARTS\b|DISTRIBUID|AUTOPECA|ROLAMENTO|MULTIMARCA|"
                     r"COMERCIO DE MAQUINAS|MAQUINAS E PECA|RETIFICA|TRATORES|DIESEL|HIDRAULIC|IMPORTA")
EMAIL_ERRADO = re.compile(r"@(outook|gmial|gmai|hotmial|hotmal|yaho)\.|@gmail\.com\.br$|\.con$")
# categorias que não entram no folder (preço subiu) — peças dessas categorias não levam 25%
FORA = re.compile(r"BRACO|LEVANTE|QUADRO|MASTRO|FILTROS? HIDRAULIC|PNEU")


def _ler(arq):
    if arq.suffix == ".tsv":
        return pd.read_csv(arq, sep="\t", dtype=str, keep_default_na=False)
    return pd.read_excel(arq, dtype=str).fillna("")


def ler(pasta):
    rep, nov = [], []
    for arq in sorted(Path(pasta).glob("*/*")):
        if arq.suffix not in (".xlsx", ".tsv"):
            continue
        d = _ler(arq)
        d["Consultor"] = arq.parent.name
        (rep if arq.name.startswith("Repeticao") else nov).append(d)
    rep, nov = pd.concat(rep, ignore_index=True), pd.concat(nov, ignore_index=True)
    for d in (rep, nov):
        d["doc"] = d["CNPJ/CPF"].map(norm.so_digitos)
        d["nome_norm"] = d.Cliente.map(norm.razao)
        d["tel"] = d.Telefone.str.strip()
        d["email"] = d["Contato (e-mail)"].str.strip().str.lower()
    rep["valor"] = pd.to_numeric(rep["Valor 1º sem/2026 (R$)"], errors="coerce")
    rep["ult"] = pd.to_datetime(rep["Data da última compra"], errors="coerce")
    nov["et"] = pd.to_datetime(nov["Data da compra da máquina (entrega técnica)"], errors="coerce")
    return rep, nov


def doc_valido(d):
    def dv(nums, pesos):
        r = sum(int(n) * p for n, p in zip(nums, pesos)) % 11
        return "0" if r < 2 else str(11 - r)
    if len(d) == 11 and len(set(d)) > 1:
        return d[9] == dv(d[:9], range(10, 1, -1)) and d[10] == dv(d[:10], range(11, 1, -1))
    if len(d) == 14 and len(set(d)) > 1:
        p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
        return d[12] == dv(d[:12], p1) and d[13] == dv(d[:13], [6] + p1)
    return False


def consultor_esperado(uf, cidade, tel):
    uf = _sem_acento(uf)
    if uf == "ES":
        return {"YURI"} if _sem_acento(cidade) in ES_NORTE else {"JHONY"}
    if uf == "MG":
        ddd = norm.so_digitos(tel)[2:4]
        return {MG_DDD[ddd], "RAYLANDER"} if ddd in MG_DDD else {"MATEUS", "DANIEL", "RAYLANDER"}
    return set(UF_MAPA.get(uf, []))


def auditar(rep, nov):
    ach = []

    def add(camp, r, item, problema, detalhe=""):
        ach.append({"Campanha": camp, "Consultor": r["Consultor"], "Cliente": r["Cliente"],
                    "CNPJ/CPF": r["doc"], "UF": r["UF"], "Item": item, "Problema": problema, "Detalhe": detalhe})

    todos = pd.concat([rep.assign(Campanha="Repetição"), nov.assign(Campanha="Novos")], ignore_index=True)
    internos = set(todos.loc[todos.email.str.contains("bmchyundai"), "tel"]) - {""}
    tel_docs = todos[todos.tel != ""].groupby("tel").doc.nunique()
    email_docs = todos[todos.email != ""].groupby("email").doc.nunique()
    docs_rep, docs_nov = set(rep.doc), set(nov.doc)
    raiz_rep = {norm.raiz(d) for d in docs_rep}

    for r in todos.to_dict("records"):
        c, d = r["Campanha"], r["doc"]
        if not doc_valido(d):
            add(c, r, "01", "CNPJ/CPF inválido (dígito verificador)", r["CNPJ/CPF"])
        if d[:8] in RAIZ_EXCLUIR and len(d) == 14:
            add(c, r, "02", "CNPJ da BMC / HD Hyundai")
        esp = consultor_esperado(r["UF"], r["Cidade"], r["tel"])
        if r["Consultor"] not in esp:
            add(c, r, "03", "Consultor fora do mapa de carteira por UF",
                f"esperado: {' ou '.join(sorted(esp)) or 'UF sem consultor'}"
                + (" (Repetição aceita o vendedor da última compra: confirmar na base)" if c == "Repetição" else ""))
        m = REVENDA.search(r["nome_norm"])
        if m:
            add(c, r, "04", "Possível revenda/oficina pelo nome", m.group(0))
        t = r["tel"]
        if not t:
            add(c, r, "05", "Sem telefone (Léo não consegue disparar)")
        elif not re.fullmatch(r"\+55[1-9][1-9]9\d{8}", t):
            add(c, r, "05", "Telefone não é celular válido (+55 DDD 9XXXX-XXXX)", t)
        if t and (t in internos or tel_docs.get(t, 0) >= 3):
            add(c, r, "06", "Telefone de consultor/BMC ou repetido em vários clientes",
                f"{t} em {tel_docs.get(t, 0)} CNPJs")
        e = r["email"]
        if "bmchyundai" in e:
            add(c, r, "07", "E-mail de funcionário BMC no cadastro do cliente", e)
        elif e and (EMAIL_ERRADO.search(e) or "@" not in e):
            add(c, r, "07", "E-mail com erro de digitação", e)
        elif e and email_docs.get(e, 0) >= 3:
            add(c, r, "07", "Mesmo e-mail em 3+ CNPJs (contador/terceiro)", e)

    # Repetição
    for r in rep.to_dict("records"):
        if pd.isna(r["ult"]) or not (INI <= r["ult"] <= FIM):
            add("Repetição", r, "10", "Última compra fora do 1º sem/2026 (pode ter comprado no 2º sem)",
                r["Data da última compra"])
        if not r["valor"] or r["valor"] <= 0:
            add("Repetição", r, "11", "Valor do 1º sem zerado/negativo", r["Valor 1º sem/2026 (R$)"])
        cats = [x.strip() for x in r["O que comprava (categorias)"].split(",") if x.strip()]
        if cats and all(FORA.search(x) for x in cats):
            add("Repetição", r, "12", "Só comprou categorias fora do folder (o 25% não se aplica ao que ele compra)",
                r["O que comprava (categorias)"])

    # duplicidades (cliente = CNPJ/CPF completo)
    for camp, df in (("Repetição", rep), ("Novos", nov)):
        g = df.groupby("doc").Consultor.unique()
        for d, cons in g.items():
            if len(cons) > 1:
                for r in df[df.doc == d].drop_duplicates("Consultor").to_dict("records"):
                    add(camp, r, "20", "Mesmo CNPJ em mais de um consultor", ", ".join(cons))
        if camp == "Repetição":
            for d, n in df.doc.value_counts().items():
                if n > 1:
                    add(camp, df[df.doc == d].iloc[0].to_dict(), "21", "CNPJ repetido na lista", f"{n} linhas")
    for d in docs_rep & docs_nov:
        r = nov[nov.doc == d].iloc[0].to_dict()
        add("Novos", r, "22", "CNPJ está em Repetição e em Novos (se comprou peça no 1º sem, não é novo)",
            f"Repetição: {', '.join(rep[rep.doc == d].Consultor.unique())}")

    # mesmo nome / mesma raiz com documento diferente: clientes distintos pela regra, mas apontar
    for r in nov.drop_duplicates("doc").to_dict("records"):
        d = r["doc"]
        if d in docs_rep:
            continue
        if len(d) == 14 and norm.raiz(d) in raiz_rep:
            outros = rep[rep.doc.map(norm.raiz) == norm.raiz(d)]
            add("Novos", r, "23", "Filial/matriz (mesma raiz de CNPJ) já compra peça e está em Repetição",
                ", ".join(f"{x.doc} ({x.Consultor})" for x in outros.itertuples()))
        nome = r["nome_norm"]
        base = re.sub(r"^\d+ ", "", re.sub(r"^\d{2} \d{3} \d{3} ", "", nome))
        mesmos = rep[(rep.nome_norm == nome) | (rep.nome_norm.map(lambda x: re.sub(r"^\d{2} \d{3} \d{3} ", "", x)) == base)]
        mesmos = mesmos[mesmos.doc != d]
        if len(base) >= 6 and len(mesmos):
            add("Novos", r, "24", "Mesmo nome com outro CNPJ/CPF em Repetição (conta separado pela regra; confirmar se não é o mesmo comprador)",
                ", ".join(f"{x.doc} ({x.Consultor})" for x in mesmos.itertuples()))
        if r["tel"] and r["tel"] not in internos and tel_docs.get(r["tel"], 0) < 3 and r["tel"] in set(rep.tel):
            outros = rep[(rep.tel == r["tel"]) & (rep.doc != d)]
            if len(outros):
                add("Novos", r, "25", "Mesmo telefone de um cliente de Repetição (provável mesmo comprador de peças)",
                    ", ".join(f"{x.Cliente} {x.doc}" for x in outros.itertuples()))

    ach = pd.DataFrame(ach).sort_values(["Item", "Consultor", "Cliente"])
    resumo = pd.concat([
        rep.groupby("Consultor").agg(**{"Repetição (CNPJs)": ("doc", "nunique"), "Repetição R$ 1º sem": ("valor", "sum")}),
        nov.groupby("Consultor").agg(**{"Novos (CNPJs)": ("doc", "nunique"), "Novos (máquinas)": ("Chassi", "nunique")}),
    ], axis=1).fillna(0)
    resumo.loc["TOTAL"] = [rep.doc.nunique(), rep.valor.sum(), nov.doc.nunique(), nov.Chassi.nunique()]
    por_item = ach.groupby(["Item", "Problema"]).agg(Ocorrencias=("Cliente", "size"),
                                                     CNPJs=("CNPJ/CPF", "nunique")).reset_index()
    return ach, resumo, por_item


def salvar(destino, **abas):
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(destino, engine="openpyxl") as w:
        for nome, df in abas.items():
            df.to_excel(w, sheet_name=nome, index=(nome == "Resumo"))
            ws = w.sheets[nome]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i in range(1, ws.max_column + 1):
                larg = max(len(str(c.value or "")) for c in ws[get_column_letter(i)][:300])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, larg + 2), 70)


if __name__ == "__main__":
    rep, nov = ler(sys.argv[1])
    ach, resumo, por_item = auditar(rep, nov)
    destino = sys.argv[2] if len(sys.argv) > 2 else "auditoria_campanha_gpd.xlsx"
    salvar(destino, Resumo=resumo, Por_item=por_item, Achados=ach)
    print(resumo.to_string())
    print(por_item.to_string(index=False))
